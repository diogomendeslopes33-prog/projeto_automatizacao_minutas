# relatorio.py
# Gera o documento Word de relatório de conformidade CCP.
# Produzido sempre — com ou sem alertas.
# Segundo documento entregue ao técnico após a minuta.

import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from utils import formatar_valor_euros
from logger import log_tecnico


# ---------------------------------------------------------------------------
# Cores
# ---------------------------------------------------------------------------

COR_ERRO        = RGBColor(0xC0, 0x00, 0x00)  # vermelho escuro
COR_AVISO       = RGBColor(0xFF, 0x7F, 0x00)  # laranja
COR_INFO        = RGBColor(0x00, 0x56, 0x96)  # azul
COR_TITULO      = RGBColor(0x1F, 0x4E, 0x79)  # azul escuro UA
COR_CINZA       = RGBColor(0x59, 0x59, 0x59)  # cinza texto secundário
COR_AVISO_CAMPO = RGBColor(0x84, 0x36, 0x00)  # castanho para [a preencher]


# ---------------------------------------------------------------------------
# Campos por preencher — lista fixa baseada na arquitectura do sistema
# ---------------------------------------------------------------------------

_DESCRICAO_CAMPOS = {
    "NR_PROCEDIMENTO":        "Número do procedimento — gerado após validação das peças",
    "NR_COMPROMISSO":         "Número do compromisso — gerado após cabimentação",
    "CLASS_ECONOMICA":        "Classificação económica — requer leitura do cabimento",
    "PRESIDENTE_JURI_CARGO":  "Cargo do presidente do júri",
    "VOGAL_EFETIVO_1_CARGO":  "Cargo do 1.º vogal efectivo",
    "VOGAL_EFETIVO_2_CARGO":  "Cargo do 2.º vogal efectivo",
    "VOGAL_SUPLENTE_1_CARGO": "Cargo do 1.º vogal suplente",
    "VOGAL_SUPLENTE_2_CARGO": "Cargo do 2.º vogal suplente",
}

_CAMPOS_EM_DESENVOLVIMENTO = {
    "ACUM_EMP_ANO_ATUAL":     "Acumulado com empresas convidadas — ano corrente (Camada B/C)",
    "ACUM_EMP_ANO_1":         "Acumulado com empresas convidadas — ano anterior (Camada B/C)",
    "ACUM_EMP_ANO_2":         "Acumulado com empresas convidadas — há dois anos (Camada B/C)",
    "ACUM_EMP_TOTAL":         "Acumulado total com empresas convidadas (Camada B/C)",
    "ACUM_EMP_REL_ANO_ATUAL": "Acumulado com empresas relacionadas — ano corrente (Camada B/C)",
    "ACUM_EMP_REL_ANO_1":     "Acumulado com empresas relacionadas — ano anterior (Camada B/C)",
    "ACUM_EMP_REL_ANO_2":     "Acumulado com empresas relacionadas — há dois anos (Camada B/C)",
    "ACUM_EMP_REL_TOTAL":     "Acumulado total com empresas relacionadas (Camada B/C)",
    "ACUM_CPV":               "Acumulado por CPV nos últimos 12 meses (Camada C)",
    "LOTE_ACUM_CPV":          "Acumulado por CPV de cada lote (Camada C)",
}


# ---------------------------------------------------------------------------
# Auxiliares de formatação
# ---------------------------------------------------------------------------

def _adicionar_titulo(doc, texto, nivel=1):
    p = doc.add_heading(texto, level=nivel)
    for run in p.runs:
        run.font.color.rgb = COR_TITULO
        run.font.name = "Arial"
    return p


def _adicionar_alerta(doc, alerta):
    """Adiciona um alerta formatado — sem prefixo de nível (redundante dentro da secção)."""
    artigo   = alerta.get("artigo", "—")
    campo    = alerta.get("campo", "")
    mensagem = alerta.get("mensagem", "")
    nivel    = alerta.get("nivel", "info").upper()

    cor = {
        "ERRO":  COR_ERRO,
        "AVISO": COR_AVISO,
        "INFO":  COR_INFO,
    }.get(nivel, COR_INFO)

    # Linha de cabeçalho — artigo + campo
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(2)

    run_artigo = p.add_run(artigo)
    run_artigo.bold = True
    run_artigo.font.name = "Arial"
    run_artigo.font.size = Pt(10)
    run_artigo.font.color.rgb = cor

    if campo:
        run_campo = p.add_run(f" | {campo}")
        run_campo.font.color.rgb = COR_CINZA
        run_campo.font.name = "Arial"
        run_campo.font.size = Pt(10)

    # Mensagem
    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Cm(0.5)
    p2.paragraph_format.space_after = Pt(4)
    run_msg = p2.add_run(mensagem)
    run_msg.font.name = "Arial"
    run_msg.font.size = Pt(9.5)


def _adicionar_linha_separadora(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    p.add_run("─" * 80).font.color.rgb = COR_CINZA


def _secao_vazia(doc, mensagem_vazia):
    p = doc.add_paragraph()
    run = p.add_run(mensagem_vazia)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = COR_CINZA
    run.italic = True


# ---------------------------------------------------------------------------
# FUNÇÃO PRINCIPAL
# ---------------------------------------------------------------------------

def gerar_relatorio(procedimento_id, dados, resultado_conformidade,
                    caminho_output, procedimento_id_log=None):
    """
    Gera o relatório de conformidade CCP em formato Word.

    Parâmetros
    ----------
    procedimento_id       : str  — identificador do procedimento (UA-AAAA-NNNN)
    dados                 : dict — campos extraídos do formulário
    resultado_conformidade: dict — output de verificar_conformidade()
    caminho_output        : str  — caminho do ficheiro .docx a gerar
    procedimento_id_log   : str  — id para o logger
    """
    log_tecnico("relatorio_inicio", {"output": caminho_output},
                procedimento_id_log or procedimento_id)

    doc = Document()

    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # -----------------------------------------------------------------------
    # Cabeçalho
    # -----------------------------------------------------------------------
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_inst.add_run("UNIVERSIDADE DE AVEIRO")
    run_inst.bold = True
    run_inst.font.name = "Arial"
    run_inst.font.size = Pt(12)
    run_inst.font.color.rgb = COR_TITULO

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Serviço de Gestão de Recursos Financeiros — AAC")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = COR_CINZA

    doc.add_paragraph()

    _adicionar_titulo(doc, "RELATÓRIO DE VERIFICAÇÃO DE CONFORMIDADE CCP", nivel=1)
    _adicionar_linha_separadora(doc)

    # Tabela de identificação
    campos_id = [
        ("Referência",         dados.get("REFERENCIA", "—")),
        ("ID do procedimento", procedimento_id),
        ("Tipo",               f"{dados.get('TIPO_PROCEDIMENTO', '—')} — {dados.get('TIPO_OBJETO', '—')}"),
        ("Objecto",            dados.get("OBJETO", "—")),
        ("Preço base",         formatar_valor_euros(dados.get("PRECO_BASE")) if dados.get("PRECO_BASE") else "—"),
        ("Data",               datetime.now().strftime("%d de %B de %Y")),
    ]

    tabela_id = doc.add_table(rows=len(campos_id), cols=2)
    tabela_id.style = "Table Grid"

    for i, (chave, valor) in enumerate(campos_id):
        celula_chave = tabela_id.rows[i].cells[0]
        celula_valor = tabela_id.rows[i].cells[1]
        celula_chave.width = Cm(5)

        p_chave = celula_chave.paragraphs[0]
        run_chave = p_chave.add_run(chave)
        run_chave.bold = True
        run_chave.font.name = "Arial"
        run_chave.font.size = Pt(9.5)

        p_valor = celula_valor.paragraphs[0]
        run_valor = p_valor.add_run(str(valor))
        run_valor.font.name = "Arial"
        run_valor.font.size = Pt(9.5)

    doc.add_paragraph()

    # -----------------------------------------------------------------------
    # Sumário
    # -----------------------------------------------------------------------
    _adicionar_titulo(doc, "Sumário", nivel=2)

    alertas      = resultado_conformidade.get("alertas", [])
    total_erros  = resultado_conformidade.get("total_erros", 0)
    total_avisos = resultado_conformidade.get("total_avisos", 0)
    total_info   = resultado_conformidade.get("total_info", 0)

    p_sum = doc.add_paragraph()
    for label, total, cor in [
        ("Erros: ",  total_erros,  COR_ERRO),
        ("Avisos: ", total_avisos, COR_AVISO),
        ("Info: ",   total_info,   COR_INFO),
    ]:
        run_l = p_sum.add_run(label)
        run_l.bold = True
        run_l.font.name = "Arial"
        run_l.font.size = Pt(10)
        run_v = p_sum.add_run(str(total) + "    ")
        run_v.bold = True
        run_v.font.color.rgb = cor
        run_v.font.name = "Arial"
        run_v.font.size = Pt(10)

    doc.add_paragraph()

    # -----------------------------------------------------------------------
    # Secção 1 — Erros
    # -----------------------------------------------------------------------
    _adicionar_titulo(doc, "1. Erros de conformidade", nivel=2)

    erros = [a for a in alertas if a.get("nivel") == "erro"]
    if erros:
        for alerta in erros:
            _adicionar_alerta(doc, alerta)
    else:
        _secao_vazia(doc, "Não foram detectados erros de conformidade.")

    doc.add_paragraph()

    # -----------------------------------------------------------------------
    # Secção 2 — Avisos
    # -----------------------------------------------------------------------
    _adicionar_titulo(doc, "2. Avisos", nivel=2)

    avisos = [a for a in alertas
              if a.get("nivel") == "aviso" and not a.get("requer_analise_juridica")]
    if avisos:
        for alerta in avisos:
            _adicionar_alerta(doc, alerta)
    else:
        _secao_vazia(doc, "Não foram detectados avisos.")

    doc.add_paragraph()

    # -----------------------------------------------------------------------
    # Secção 3 — Informações
    # -----------------------------------------------------------------------
    _adicionar_titulo(doc, "3. Informações", nivel=2)

    infos = [a for a in alertas
             if a.get("nivel") == "info" and not a.get("requer_analise_juridica")]
    if infos:
        for alerta in infos:
            _adicionar_alerta(doc, alerta)
    else:
        _secao_vazia(doc, "Sem informações adicionais.")

    doc.add_paragraph()

    # -----------------------------------------------------------------------
    # Secção 4 — Análise jurídica
    # -----------------------------------------------------------------------
    _adicionar_titulo(doc, "4. Requer análise", nivel=2)

    juridicas = [a for a in alertas if a.get("requer_analise_juridica")]
    if juridicas:
        for alerta in juridicas:
            _adicionar_alerta(doc, alerta)
    else:
        _secao_vazia(doc, "Nenhuma norma interpretativa aplicável a este procedimento.")

    doc.add_paragraph()

    # -----------------------------------------------------------------------
    # Secção 5 — Campos por preencher manualmente
    # -----------------------------------------------------------------------
    _adicionar_titulo(doc, "5. Campos que requerem preenchimento manual", nivel=2)

    p_intro = doc.add_paragraph()
    run_intro = p_intro.add_run(
        "Os seguintes campos não foram preenchidos automaticamente "
        "e requerem intervenção manual do técnico antes de submeter as peças:"
    )
    run_intro.font.name = "Arial"
    run_intro.font.size = Pt(10)

    for campo, descricao in _DESCRICAO_CAMPOS.items():
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.5)
        run_campo = p.add_run(f"{campo}  ")
        run_campo.bold = True
        run_campo.font.name = "Arial"
        run_campo.font.size = Pt(10)
        run_campo.font.color.rgb = COR_AVISO_CAMPO
        run_desc = p.add_run(f"— {descricao}")
        run_desc.font.name = "Arial"
        run_desc.font.size = Pt(10)

    doc.add_paragraph()

    # -----------------------------------------------------------------------
    # Secção 6 — Campos em desenvolvimento (Camadas B e C)
    # -----------------------------------------------------------------------
    _adicionar_titulo(doc, "6. Campos que aguardam implementação (Camadas B e C)", nivel=2)

    p_dev_intro = doc.add_paragraph()
    run_dev_intro = p_dev_intro.add_run(
        "Os seguintes campos serão preenchidos automaticamente quando as "
        "Camadas B e C estiverem activas. Ficam temporariamente marcados "
        "como [em desenvolvimento] na minuta:"
    )
    run_dev_intro.font.name = "Arial"
    run_dev_intro.font.size = Pt(10)
    run_dev_intro.font.color.rgb = COR_CINZA
    run_dev_intro.italic = True

    for campo, descricao in _CAMPOS_EM_DESENVOLVIMENTO.items():
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.5)
        run_campo = p.add_run(f"{campo}  ")
        run_campo.bold = True
        run_campo.font.name = "Arial"
        run_campo.font.size = Pt(10)
        run_campo.font.color.rgb = COR_CINZA
        run_desc = p.add_run(f"— {descricao}")
        run_desc.font.name = "Arial"
        run_desc.font.size = Pt(10)
        run_desc.font.color.rgb = COR_CINZA

    doc.add_paragraph()
    _adicionar_linha_separadora(doc)

    # Rodapé
    p_rod = doc.add_paragraph()
    p_rod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_rod = p_rod.add_run(
        f"Relatório gerado automaticamente pelo Sistema de Automação de Contratação Pública — UA"
        f"  |  {procedimento_id}"
    )
    run_rod.font.name = "Arial"
    run_rod.font.size = Pt(8)
    run_rod.font.color.rgb = COR_CINZA

    # Guarda
    os.makedirs(os.path.dirname(caminho_output), exist_ok=True)
    doc.save(caminho_output)

    log_tecnico("relatorio_gerado", {"output": caminho_output},
                procedimento_id_log or procedimento_id)

    return caminho_output