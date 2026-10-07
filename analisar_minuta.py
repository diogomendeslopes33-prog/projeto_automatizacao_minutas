# analisar_minuta.py
# Verificação 2 (double-check) — extrai da minuta FINAL GERADA (não do
# formulário original) os campos que o técnico pode ter corrigido
# manualmente depois da geração, e volta a correr verificar_conformidade()
# sobre eles. Compara com os alertas da Verificação 1 (via
# historico_procedimento) e calcula o delta: resolvidos / novos / persistentes.
#
# ÂMBITO (Decisão de arquitectura desta sessão): só se re-extraem os
# campos que o técnico pode razoavelmente corrigir à mão na minuta Word
# depois de ver um alerta no relatório — Caução, Lotes, Prazo Contratual,
# Júri e Gestor/Substituto do Contrato. NÃO se re-extraem:
#   - EMPRESAS_CONVIDADAS (relacionadas, acum_fornecedor), ACUM_CPV,
#     ACUM_OBJETO — vêm de fontes externas (ERP, Informa D&B); editar o
#     número na minuta não muda o valor real que esse número representa.
#     Estes são reutilizados tal como estavam na Verificação 1.
#   - TIPO_PROCEDIMENTO, TIPO_OBJETO, PRECO_BASE — mudar estes implicaria
#     mudar de minuta inteira, não uma correcção pontual.
#   - DESPESA_PLURIANUAL, GARANTIA_ANOS — não existem como cláusula na
#     minuta "Consulta Prévia + Bens" actualmente construída; a
#     re-extracção destes campos só se aplica a minutas que os incluam.
#
# DUAS FORMAS DE CAMPO NA MINUTA:
#   1. Blocos condicionais (Caução, Lotes) — o marcador de controlo
#      {{SE_X}}/{{FIM_X}} já foi substituído: só um dos dois ramos de
#      texto sobrevive. Detecta-se por PRESENÇA de uma frase distintiva
#      de cada ramo, não por leitura de um valor numa célula.
#   2. Marcadores simples embutidos em frase (Prazo, Júri, Gestor) — o
#      valor está onde o marcador estava; extrai-se por regex ancorado
#      no texto literal que rodeia o marcador no molde.
#
# CONSISTÊNCIA: campos que aparecem mais do que uma vez na minuta (ex:
# Caução, Lotes — repetidos em anexos diferentes) são verificados em
# TODAS as ocorrências. Se as frases dos dois ramos aparecerem ambas no
# documento, é sinal de edição manual inconsistente entre ocorrências —
# gera o seu próprio alerta, em vez de escolher uma arbitrariamente.

import re
from docx import Document

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _texto_documento_completo(doc):
    """
    Texto de todos os parágrafos e tabelas, lido nó a nó (<w:t>) — evita
    perder texto fragmentado entre múltiplos runs, o mesmo cuidado já
    tomado em preencher_word.py e apinforma.py ao longo do projecto.
    """
    partes = []
    for p in doc.paragraphs:
        partes.append("".join(t.text or "" for t in p._element.iter(f"{{{W_NS}}}t")))
    for tabela in doc.tables:
        for linha in tabela.rows:
            partes.append("".join(t.text or "" for t in linha._tr.iter(f"{{{W_NS}}}t")))
    return "\n".join(partes)


def _detectar_valor_binario(texto_completo, frase_sim, frase_nao, nome_campo):
    """
    Detecta qual dos dois ramos de um bloco condicional sobreviveu na
    minuta, procurando a frase distintiva de cada ramo em todo o texto.

    Devolve (valor, inconsistente):
        valor         — "sim", "nao", ou None se nenhuma frase encontrada
        inconsistente — True se AMBAS as frases foram encontradas (sinal
                        de edição manual inconsistente entre ocorrências
                        repetidas do mesmo campo condicional na minuta)
    """
    tem_sim = frase_sim in texto_completo
    tem_nao = frase_nao in texto_completo

    if tem_sim and tem_nao:
        return None, True
    if tem_sim:
        return "sim", False
    if tem_nao:
        return "nao", False
    return None, False


def extrair_dados_minuta(caminho_docx):
    """
    Extrai da minuta final gerada os campos relevantes para conformidade,
    dentro do âmbito definido (ver cabeçalho do módulo).

    Devolve dict:
        {
            "CAUCAO": "Não" | "Sim" | None,
            "CAUCAO_inconsistente": bool,
            "TEM_LOTES": False | True | None,
            "TEM_LOTES_inconsistente": bool,
            "PRAZO_ENTREGA": str | None,
            "PRAZO_ENTREGA_UNIDADE": str | None,
            "PRESIDENTE_JURI_NOME": str | None,
            "VOGAL_EFETIVO_1_NOME": str | None,
            "VOGAL_EFETIVO_2_NOME": str | None,
            "VOGAL_SUPLENTE_1_NOME": str | None,
            "VOGAL_SUPLENTE_2_NOME": str | None,
            "GESTOR_CONTRATO": str | None,
            "SUBSTITUTO_GESTOR_CONTRATO": str | None,
        }
    Devolve None se o ficheiro não puder ser aberto.
    """
    try:
        doc = Document(caminho_docx)
    except Exception:
        return None

    texto = _texto_documento_completo(doc)
    dados = {}

    # --- CAUÇÃO (bloco condicional) ---
    frase_caucao_sim = (
        "Pretende-se a apresentação de caução à luz do n.º 1 do artigo 88.º do CCP"
    )
    frase_caucao_nao = (
        "Foi dispensada a apresentação de caução nos termos da alínea a) "
        "do n.º 2 do artigo 88.º do CCP"
    )
    valor_caucao, inconsist_caucao = _detectar_valor_binario(
        texto, frase_caucao_sim, frase_caucao_nao, "CAUCAO"
    )
    dados["CAUCAO"] = {"sim": "Sim", "nao": "Não"}.get(valor_caucao)
    dados["CAUCAO_inconsistente"] = inconsist_caucao

    # --- LOTES (bloco condicional) ---
    frase_com_lotes = "Os números de cabimento da presente despesa são os seguintes"
    frase_sem_lotes = "O número de cabimento da presente despesa é"
    valor_lotes, inconsist_lotes = _detectar_valor_binario(
        texto, frase_com_lotes, frase_sem_lotes, "TEM_LOTES"
    )
    dados["TEM_LOTES"] = {"sim": True, "nao": False}.get(valor_lotes)
    dados["TEM_LOTES_inconsistente"] = inconsist_lotes

    # --- PRAZO CONTRATUAL (marcador simples) ---
    m = re.search(
        r"O prazo contratual do presente procedimento é de\s+(\S+)\s+(.+?),\s*"
        r"contados desde a data da aposição",
        texto
    )
    dados["PRAZO_ENTREGA"] = m.group(1).strip() if m else None
    dados["PRAZO_ENTREGA_UNIDADE"] = m.group(2).strip() if m else None

    # --- JÚRI (marcadores simples, um por linha) ---
    padroes_juri = {
        "PRESIDENTE_JURI_NOME":   r"Presidente:\s*(.+?)\s*,",
        "VOGAL_EFETIVO_1_NOME":   r"Primeiro Vogal Efetivo:\s*(.+?)\s*,",
        "VOGAL_EFETIVO_2_NOME":   r"Segundo Vogal Efetivo:\s*(.+?)\s*,",
        "VOGAL_SUPLENTE_1_NOME":  r"Primeiro Vogal Suplente:\s*(.+?)\s*,",
        "VOGAL_SUPLENTE_2_NOME":  r"Segundo Vogal Suplente:\s*(.+?)\s*,",
    }
    for campo, padrao in padroes_juri.items():
        m = re.search(padrao, texto)
        dados[campo] = m.group(1).strip() if m else None

    # --- GESTOR DO CONTRATO E SUBSTITUTO (marcadores simples, mesma frase) ---
    # Ancorado em "no artigo 290º-A do CCP," (texto fixo do molde) em vez
    # da primeira vírgula a seguir a "gestor do contrato" — a frase tem
    # uma cláusula legal intercalada antes do nome, com a sua própria
    # vírgula, o que fazia o regex anterior capturar texto a mais.
    m = re.search(
        r"no artigo 290º-A do CCP,\s*(.+?),\s*a exercer funções de.*?"
        r"desta,\s*(.+?),\s*a exercer funções",
        texto
    )
    # Nomenclatura conformidade/ (GESTOR_CONTRATO), não a de ler_pdf.py
    # (GESTOR_NOME) — ver normalização já feita em main.py (Decisão 36)
    dados["GESTOR_CONTRATO"] = m.group(1).strip() if m else None
    dados["SUBSTITUTO_GESTOR_CONTRATO"] = m.group(2).strip() if m else None

    return dados


def _chave_alerta(alerta):
    """
    Chave de identidade de um alerta, para comparação entre verificações.
    Usa (artigo, campo, mensagem) — dois alertas são "o mesmo" se
    coincidirem nestes três valores.
    """
    return (alerta.get("artigo"), alerta.get("campo"), alerta.get("mensagem"))


def comparar_com_verificacao_1(procedimento_id, caminho_minuta):
    """
    Verificação 2 — orquestrador completo. Pode ser chamado numa
    invocação do programa totalmente separada da que gerou a minuta
    (Verificação 1) — não depende de nenhum estado em memória; tudo o
    que precisa vem do sequencia/ e da própria minuta.

    1. Recupera de historico_procedimento() os alertas da Verificação 1
       e os "dados_reutilizaveis" (campos que NÃO são re-extraídos —
       EMPRESAS_CONVIDADAS, ACUM_CPV, ACUM_OBJETO, TIPO_PROCEDIMENTO,
       TIPO_OBJETO, PRECO_BASE — ver "ÂMBITO" no cabeçalho do módulo).
    2. Re-extrai da minuta final gerada os campos re-verificáveis
       (Caução, Lotes, Prazo, Júri, Gestor/Substituto).
    3. Junta os dois conjuntos e corre verificar_conformidade().
    4. Calcula o delta: resolvidos / novos / persistentes.

    Devolve dict:
        {
            "dados_reextraidos": {...},
            "campos_inconsistentes": [...],
            "alertas_v1": [...],
            "alertas_v2": [...],
            "resolvidos":    [...],
            "novos":         [...],
            "persistentes":  [...],
        }
    Devolve None se o procedimento não existir no sequencia/, se a
    Verificação 1 não tiver dados detalhados guardados (procedimentos
    antigos, anteriores a esta correcção), ou se a minuta não puder ser lida.
    """
    from sequencia.sequencia import historico_procedimento
    from conformidade import verificar_conformidade

    historico = historico_procedimento(procedimento_id)
    if historico is None:
        return None

    evento_v1 = next(
        (e for e in historico["eventos"] if e["evento"] == "conformidade_verificada"),
        None
    )
    if (evento_v1 is None or not evento_v1.get("dados")
            or "alertas" not in evento_v1["dados"]
            or "dados_reutilizaveis" not in evento_v1["dados"]):
        return None

    alertas_v1 = evento_v1["dados"]["alertas"]
    dados_reutilizaveis = evento_v1["dados"]["dados_reutilizaveis"]

    dados_minuta = extrair_dados_minuta(caminho_minuta)
    if dados_minuta is None:
        return None

    campos_inconsistentes = [
        campo.replace("_inconsistente", "")
        for campo, valor in dados_minuta.items()
        if campo.endswith("_inconsistente") and valor
    ]

    # Dados para a Verificação 2: reutilizáveis da V1 + re-extraídos da minuta
    dados_v2 = dict(dados_reutilizaveis)
    for campo, valor in dados_minuta.items():
        if not campo.endswith("_inconsistente"):
            dados_v2[campo] = valor

    resultado_v2 = verificar_conformidade(dados_v2)
    alertas_v2 = resultado_v2["alertas"]

    chaves_v1 = {_chave_alerta(a): a for a in alertas_v1}
    chaves_v2 = {_chave_alerta(a): a for a in alertas_v2}

    resolvidos = [a for k, a in chaves_v1.items() if k not in chaves_v2]
    novos      = [a for k, a in chaves_v2.items() if k not in chaves_v1]
    persistentes = [a for k, a in chaves_v2.items() if k in chaves_v1]

    return {
        "dados_reextraidos":     dados_minuta,
        "campos_inconsistentes": campos_inconsistentes,
        "alertas_v1":            alertas_v1,
        "alertas_v2":            alertas_v2,
        "resolvidos":            resolvidos,
        "novos":                 novos,
        "persistentes":          persistentes,
    }


def gerar_relatorio_verificacao2(procedimento_id, resultado, caminho_output):
    """
    Gera um documento Word simples com o resultado da Verificação 2 —
    mesmo espírito do relatório de conformidade da Verificação 1
    (relatorio.py): identificação mínima + secções, sem reproduzir dados
    que já constam noutro lado. Secções vazias mostram frase neutra.
    """
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading(f"Verificação 2 — {procedimento_id}", level=1)
    doc.add_paragraph(
        "Double-check sobre a minuta final gerada, comparando com os "
        "alertas de conformidade da Verificação 1 (sobre os dados do "
        "formulário)."
    )

    if resultado["campos_inconsistentes"]:
        doc.add_heading("Inconsistências detectadas", level=2)
        doc.add_paragraph(
            "Os seguintes campos aparecem com valores contraditórios em "
            "diferentes partes da minuta — confirmar manualmente qual é "
            "o valor correcto:"
        )
        for campo in resultado["campos_inconsistentes"]:
            doc.add_paragraph(f"• {campo}", style="List Bullet")

    secoes = [
        ("Erros/avisos resolvidos", resultado["resolvidos"],
         "Não foram identificados alertas resolvidos face à Verificação 1."),
        ("Erros/avisos novos", resultado["novos"],
         "Não foram identificados novos alertas face à Verificação 1."),
        ("Erros/avisos persistentes", resultado["persistentes"],
         "Não existem alertas persistentes por resolver."),
    ]

    for titulo, alertas, frase_vazia in secoes:
        doc.add_heading(titulo, level=2)
        if not alertas:
            doc.add_paragraph(frase_vazia)
            continue
        for alerta in alertas:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"[{alerta.get('nivel', '').upper()}] ").bold = True
            p.add_run(f"{alerta.get('artigo', '')} — {alerta.get('mensagem', '')}")

    doc.save(caminho_output)
    return caminho_output